"""
合并 output/ 下所有 xlsx，按平台分 sheet，跨文件去重（保留最新抓取时间的条目）。
用法：python scripts/dedup_output.py [--output-dir <dir>] [--out <file>]
"""
import argparse
import glob
import os
import sys
from datetime import datetime

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def parse_time(val) -> datetime:
    if isinstance(val, datetime):
        return val
    if val is None:
        return datetime.min
    s = str(val).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y/%m/%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    return datetime.min


def load_sheet(path: str) -> tuple[list, list[list]]:
    """返回 (header, rows)，rows 是 list of list（values_only）。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    all_rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    if not all_rows:
        return [], []
    return all_rows[0], all_rows[1:]


def dedup(files: list[str], id_col_names: tuple[str, ...], time_col_name: str):
    """
    读取多个 xlsx，以 id_col_names 之一为去重键，保留抓取时间最新的行。
    返回 (header, deduped_rows)。
    """
    header: list | None = None
    id_col: int = -1
    time_col: int = -1
    # id -> (row, time)
    best: dict[str, tuple[list, datetime]] = {}

    for f in sorted(files):
        h, rows = load_sheet(f)
        if not h:
            continue
        if header is None:
            header = h
            id_col = next((i for i, c in enumerate(h) if c in id_col_names), -1)
            time_col = next((i for i, c in enumerate(h) if c == time_col_name), -1)
        for row in rows:
            if id_col < 0 or id_col >= len(row):
                continue
            rid = str(row[id_col]).strip() if row[id_col] is not None else ""
            if not rid:
                continue
            t = parse_time(row[time_col] if 0 <= time_col < len(row) else None)
            if rid not in best or t > best[rid][1]:
                best[rid] = (row, t)

    deduped = [v[0] for v in best.values()]
    return header or [], deduped


def write_sheet(ws, header: list, rows: list[list]):
    # 写表头
    ws.append(header)
    header_row = ws[1]
    fill = PatternFill("solid", fgColor="1F4E79")
    for cell in header_row:
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 写数据
    for row in rows:
        ws.append(row)

    # 自动列宽（采样）
    for col_idx, _ in enumerate(header, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(header[col_idx - 1] or ""))
        for row in rows[:200]:
            val = row[col_idx - 1] if col_idx - 1 < len(row) else ""
            max_len = max(max_len, len(str(val or "")))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 40)

    ws.freeze_panes = "A2"


def main():
    parser = argparse.ArgumentParser(description="合并去重 output/ 目录下的抓取结果")
    parser.add_argument("--output-dir", default="output", help="xlsx 所在目录")
    parser.add_argument("--out", default=None, help="输出文件路径（默认 output/merged-去重-<时间>.xlsx）")
    args = parser.parse_args()

    out_dir = args.output_dir
    if not os.path.isdir(out_dir):
        print(f"[ERROR] 目录不存在: {out_dir}", file=sys.stderr)
        sys.exit(1)

    douyin_files = sorted(glob.glob(os.path.join(out_dir, "douyin-*.xlsx")))
    xhs_files = sorted(glob.glob(os.path.join(out_dir, "xhs-*.xlsx")))

    if not douyin_files and not xhs_files:
        print("[WARN] 未找到任何抓取结果文件")
        sys.exit(0)

    print(f"抖音文件: {len(douyin_files)} 个")
    print(f"小红书文件: {len(xhs_files)} 个")

    dy_header, dy_rows = dedup(douyin_files, ("图文ID",), "抓取时间")
    xhs_header, xhs_rows = dedup(xhs_files, ("笔记ID",), "抓取时间")

    print(f"抖音去重后: {len(dy_rows)} 条")
    print(f"小红书去重后: {len(xhs_rows)} 条")

    ts = datetime.now().strftime("%Y%m%d-%H%M%S")
    out_path = args.out or os.path.join(out_dir, f"merged-去重-{ts}.xlsx")

    wb = Workbook()
    wb.remove(wb.active)  # 删默认空 sheet

    if dy_header and dy_rows:
        ws_dy = wb.create_sheet("抖音")
        write_sheet(ws_dy, dy_header, dy_rows)

    if xhs_header and xhs_rows:
        ws_xhs = wb.create_sheet("小红书")
        write_sheet(ws_xhs, xhs_header, xhs_rows)

    wb.save(out_path)
    print(f"\n[OK] 已保存到: {out_path}")
    print(f"     抖音 sheet: {len(dy_rows)} 行 / 小红书 sheet: {len(xhs_rows)} 行")


if __name__ == "__main__":
    main()
